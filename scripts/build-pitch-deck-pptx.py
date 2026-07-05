#!/usr/bin/env python3
"""Build Krunchr pitch deck PPTX from docs/pitch-deck.md content.

Run: python scripts/build-pitch-deck-pptx.py
"""
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml

ROOT = Path(__file__).parent.parent
DOCS = ROOT / "docs"
PUBLIC = ROOT / "public"
OUTPUT = DOCS / "pitch-deck-Krunchr.pptx"

# Brand colors
TEAL = RGBColor(13, 148, 136)
TEAL_INK = RGBColor(0, 106, 97)
MINT = RGBColor(78, 222, 163)
INK = RGBColor(11, 28, 48)
MUTED = RGBColor(69, 70, 77)
ERROR = RGBColor(186, 26, 26)
CANVAS = RGBColor(248, 249, 255)
CARD = RGBColor(255, 255, 255)
BORDER = RGBColor(226, 232, 240)
NAVY = RGBColor(11, 28, 48)
WHITE = RGBColor(255, 255, 255)


def hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def render_logo(filename: str, bg: str, stroke: str, node: str, size: int = 512) -> Path:
    """Render the Krunchr logo as PNG from primitive shapes."""
    img = Image.new("RGBA", (size, size), bg)
    draw = ImageDraw.Draw(img)
    r = size / 512

    # Chain-link node (rounded rect with hole)
    draw.rounded_rectangle(
        [int(336 * r), int(192 * r), int(432 * r), int(288 * r)],
        radius=int(28 * r),
        fill=node,
    )
    draw.rounded_rectangle(
        [int(360 * r), int(216 * r), int(408 * r), int(264 * r)],
        radius=int(14 * r),
        fill=bg,
    )

    # Checkmark stroke
    sw = int(56 * r)
    p1 = (int(128 * r), int(168 * r))
    p2 = (int(240 * r), int(344 * r))
    p3 = (int(336 * r), int(240 * r))
    draw.line([p1, p2, p3], fill=stroke, width=sw, joint="curve")

    # Rounded caps / joints
    cap = sw // 2
    for p in (p1, p2, p3):
        draw.ellipse([p[0] - cap, p[1] - cap, p[0] + cap, p[1] + cap], fill=stroke)

    out = DOCS / filename
    img.save(out)
    return out


def set_slide_bg(slide, color: RGBColor):
    """Set solid background color for a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    font_size: int,
    color: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font_name: str = "SF Pro Display",
):
    """Add a text box with basic styling."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return box


def add_bullet_list(slide, left, top, width, height, items, font_size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "SF Pro Text"
        p.space_after = Pt(16)
        # Use disc bullet via paragraph formatting
        p._pPr.insert(0, parse_xml(r'<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="•"/>'))
    return box


def add_card(slide, left, top, width, height, title, body, number=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    shape.adjustments[0] = 0.08

    if number is not None:
        num = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.25), Inches(top + 0.25), Inches(0.45), Inches(0.45))
        num.fill.solid()
        num.fill.fore_color.rgb = TEAL
        num.line.fill.background()
        add_textbox(slide, left + 0.25, top + 0.30, 0.45, 0.35, str(number), 18, WHITE, True, PP_ALIGN.CENTER)

    add_textbox(slide, left + 0.25 + (0.55 if number is not None else 0), top + 0.25, width - 0.5 - (0.55 if number is not None else 0), 0.4, title, 20, INK, True)
    add_textbox(slide, left + 0.25, top + 0.75, width - 0.5, height - 1.0, body, 14, MUTED)


def add_footer(slide, logo_path: Path, slide_num: int, total: int):
    # Footer line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.95), Inches(12.13), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()

    # Logo + brand
    slide.shapes.add_picture(str(logo_path), Inches(0.6), Inches(7.05), width=Inches(0.35))
    add_textbox(slide, 1.05, 7.05, 2.0, 0.3, "Krunchr", 14, INK, True)

    # Progress
    add_textbox(slide, 11.8, 7.05, 0.9, 0.3, f"{slide_num} / {total}", 14, MUTED, False, PP_ALIGN.RIGHT)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # blank

    # Render logos
    logo_colored = render_logo("logo-colored-pptx.png", "#F8F9FF", "#0D9488", "#4EDEA3")
    logo_white = render_logo("logo-white-pptx.png", "#0B1C30", "#FFFFFF", "#FFFFFF")
    logo_dark = render_logo("logo-dark-pptx.png", "#FFFFFF", "#0B1C30", "#0B1C30")

    # Common slide dimensions
    W = 13.333
    H = 7.5

    # Slide 1: Title
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    slide.shapes.add_picture(str(logo_colored), Inches(0.8), Inches(0.9), width=Inches(1.1))
    add_textbox(slide, 0.8, 2.2, 11.5, 1.2, "Krunchr", 84, INK, True)
    add_textbox(slide, 0.8, 3.4, 11.0, 0.7, "Compliance Engine — Philippine tax filing, automated.", 32, MUTED)
    add_textbox(slide, 0.8, 4.4, 10.5, 1.0, "Turns one of the most anxious chores a Filipino freelancer faces into something that just feels handled.", 26, MUTED)
    add_textbox(slide, 0.8, 5.8, 11.5, 0.4, "APAC Stellar Hackathon 2026 · Local Finance & Real World Access", 16, TEAL_INK, True)
    add_textbox(slide, 0.8, 6.2, 11.5, 0.3, "Team: [PLACEHOLDER]    Contact: krunchr@artisam.xyz", 15, MUTED)
    add_footer(slide, logo_colored, 1, 8)

    # Slide 2: Problem
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "THE PROBLEM", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.9, "Eight returns a year. One mistake costs money, time, and trust.", 48, INK, True)
    items = [
        "A freelancer on the 8% flat rate owes up to 8 BIR returns a year — 2551Q ×4, 1701Q ×3, 1701A — in strict legal order.",
        "The rules interact badly: the 8% election is irrevocable; mixed-income earners lose the ₱250,000 exemption; penalties changed under RA 11976; due dates roll past Philippine holidays.",
        "One wrong sequence or missed deadline → surcharge, daily interest, compromise penalties, and a black mark banks and embassies ask about.",
        "Once filed, all a freelancer has is a scanned PDF — no quick, trustworthy way for a third party to verify it.",
    ]
    add_bullet_list(slide, 0.8, 2.15, 7.0, 4.8, items, 19)
    # Decorative paper stack placeholder shape
    stack = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(2.0), Inches(3.8), Inches(4.2))
    stack.fill.solid()
    stack.fill.fore_color.rgb = CARD
    stack.line.color.rgb = BORDER
    add_textbox(slide, 8.7, 3.7, 3.6, 1.0, "[Generated image placeholder: overwhelmed freelancer buried in tax forms and deadlines]", 13, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide, logo_colored, 2, 8)

    # Slide 3: Solution
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "THE SOLUTION", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.9, "Krunchr turns eight confusing returns into one upload.", 48, INK, True)
    cards = [
        (0.8, 2.2, "Upload Form 2307s", "Drop in withholding certificates. The engine recomputes your entire tax year, every return, in the correct legal sequence."),
        (3.5, 2.2, "Correct by Construction", "8% rules, RA 11976 penalty rates, holiday-aware due dates, and mixed-income exemption rules are enforced in code."),
        (6.2, 2.2, "Tamper-Evident Receipt", "Every filed return is anchored on Stellar. Banks, embassies, or auditors can verify it in seconds, not trust a PDF."),
        (8.9, 2.2, "One-Click Filing Package", "All return PDFs, the SAWT summary, and a cover sheet — bundled and ready to download."),
    ]
    for i, (x, y, title, body) in enumerate(cards, 1):
        add_card(slide, x, y, 2.5, 3.8, title, body, i)
    add_footer(slide, logo_colored, 3, 8)

    # Slide 4: Demo
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "LIVE DEMO", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.6, "The single demo moment", 48, INK, True)
    add_textbox(slide, 0.8, 1.6, 11.0, 0.4, "From onboarding to verified compliance in about 30 seconds.", 24, MUTED)
    steps = [
        ("Onboard", "Maria finishes a 3-minute setup: TIN, RDO, tax codes."),
        ("Upload 2307s", "Her withholding certificates land in the dashboard."),
        ("Returns Compute", "All 8 returns run in sequence: Blocked → Pending → Generated → Filed."),
        ("File & Anchor", "A Stellar TX ID and scannable QR appear on the receipt."),
        ("Verify", "A loan officer scans it — compliance confirmed in ~30 seconds."),
    ]
    step_w = 2.2
    gap = 0.22
    start_x = 0.8
    for i, (title, body) in enumerate(steps, 1):
        x = start_x + (i - 1) * (step_w + gap)
        add_card(slide, x, 2.2, step_w, 2.6, title, body, i)
    # Demo placeholder
    demo_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.73), Inches(1.7))
    demo_box.fill.solid()
    demo_box.fill.fore_color.rgb = CARD
    demo_box.line.color.rgb = BORDER
    add_textbox(slide, 0.8, 5.65, 11.73, 0.5, "[Demo screen recording placeholder — ~60–90s video: onboarding → upload 2307 → returns compute → file one → QR / verification]", 16, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide, logo_colored, 4, 8)

    # Slide 5: How it works
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "HOW IT WORKS", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.6, "One app. One engine. One immutable receipt.", 48, INK, True)
    arch = [
        ("Upload income", "Form 2307 certificates trigger a full-year recascade."),
        ("Pure computation engine", "Every return, in order, with RA 11976 penalties and holiday deadlines."),
        ("Generate BIR PDF", "Official-form PDFs plus SAWT and cover sheet."),
        ("Anchor hash on Stellar", "A SHA-256 fingerprint is stored as a manageData operation per return."),
    ]
    y = 1.9
    for title, body in arch:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(6.5), Inches(0.85))
        card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = BORDER
        add_textbox(slide, 1.0, y + 0.12, 6.1, 0.3, title, 17, INK, True)
        add_textbox(slide, 1.0, y + 0.42, 6.1, 0.35, body, 13, MUTED)
        if y < 5.0:
            arrow = add_textbox(slide, 3.8, y + 0.82, 0.5, 0.25, "↓", 20, TEAL, True, PP_ALIGN.CENTER)
        y += 1.05
    items = [
        "Money is never floating-point — all math runs through decimal.js, so a peso is always a peso.",
        "Append-only audit trail — elections, filings, retries, and dispositions are logged forever.",
        "Failure is graceful — if anchoring hiccups, the filing still succeeds and retries later.",
    ]
    add_bullet_list(slide, 8.0, 2.0, 4.5, 4.2, items, 18)
    add_footer(slide, logo_colored, 5, 8)

    # Slide 6: Market
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "IMPACT & MARKET", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.6, "Local finance, with real-world access unlocked.", 48, INK, True)
    items = [
        "Who needs this: millions of Filipino freelancers on the 8% flat rate who file manually, hire quarterly accountants, or don't file and accumulate penalties.",
        "Why now: the BIR 8% regime, the Ease of Paying Taxes Act, and a growing freelance economy make a guided, correct filing path genuinely useful.",
        "The trust unlock: a verifiable receipt turns a private chore into something banks, embassies, and auditors confirm in seconds.",
        "Built for the track: local finance first, with real-world access via the Stellar-anchored receipt.",
    ]
    add_bullet_list(slide, 0.8, 2.0, 7.5, 4.5, items, 19)
    stats = [("8 → 1", "Returns turned into one upload"), ("~30s", "Third-party compliance verification"), ("∞", "Tamper-evident audit trail")]
    for i, (stat, label) in enumerate(stats):
        y = 2.0 + i * 1.5
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(y), Inches(3.7), Inches(1.25))
        card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = BORDER
        add_textbox(slide, 8.9, y + 0.12, 3.5, 0.6, stat, 40, TEAL, True, PP_ALIGN.CENTER)
        add_textbox(slide, 8.9, y + 0.72, 3.5, 0.4, label, 14, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide, logo_colored, 6, 8)

    # Slide 7: Roadmap
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    add_textbox(slide, 0.8, 0.6, 11.0, 0.4, "WHAT'S NEXT", 15, TEAL_INK, True)
    add_textbox(slide, 0.8, 1.0, 11.0, 0.6, "Honest gaps that form the roadmap.", 48, INK, True)
    roadmap = [
        ("Graduated-rate path", "8% is the finished demo; graduated bracket math exists but the end-to-end flow isn't complete.", "In progress"),
        ("Mixed-income → Form 1701", "Routing and the 1701 PDF template exist; the full annual-return flow is the next mile.", "In progress"),
        ("VAT-threshold enforcement", "Warning at ₱3M exists; the hard 'register for VAT' block is next.", "In progress"),
        ("Self-serve registration", "Accounts are admin-seeded today; a public sign-up flow is needed for real adoption.", "Future"),
        ("OCR for Form 2307", "Auto-read the certificate with a phone camera instead of typing figures.", "Future"),
        ("Mainnet + public verifier", "Move off testnet and give third parties a one-link verification page.", "Future"),
    ]
    positions = [
        (0.8, 2.0), (4.65, 2.0), (8.5, 2.0),
        (0.8, 4.1), (4.65, 4.1), (8.5, 4.1),
    ]
    for (x, y), (title, body, status) in zip(positions, roadmap):
        add_card(slide, x, y, 3.55, 1.8, title, body)
        status_color = TEAL_INK if status == "In progress" else MUTED
        add_textbox(slide, x + 0.25, y + 0.25, 2.0, 0.25, status.upper(), 11, status_color, True)
    add_footer(slide, logo_colored, 7, 8)

    # Slide 8: Team / Thanks
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, CANVAS)
    slide.shapes.add_picture(str(logo_colored), Inches(6.15), Inches(1.0), width=Inches(1.0))
    add_textbox(slide, 0.8, 2.2, 11.73, 0.8, "Thank you.", 60, INK, True, PP_ALIGN.CENTER)
    add_textbox(slide, 0.8, 3.0, 11.73, 0.4, "Eight returns. One upload. Handled.", 28, MUTED, False, PP_ALIGN.CENTER)
    for i in range(3):
        x = 3.1 + i * 2.4
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.8), Inches(2.0), Inches(2.0))
        card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = BORDER
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.65), Inches(4.05), Inches(0.7), Inches(0.7))
        circle.fill.solid(); circle.fill.fore_color.rgb = hex_to_rgb("#E5EEFF"); circle.line.fill.background()
        add_textbox(slide, x, 4.85, 2.0, 0.3, "[Name]", 18, INK, True, PP_ALIGN.CENTER)
        add_textbox(slide, x, 5.2, 2.0, 0.25, "[Role]", 14, MUTED, False, PP_ALIGN.CENTER)
    add_textbox(slide, 0.8, 6.1, 11.73, 0.3, "app.krunchr.xyz    ·    krunchr.xyz    ·    krunchr@artisam.xyz", 17, TEAL_INK, True, PP_ALIGN.CENTER)
    add_footer(slide, logo_colored, 8, 8)

    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    build()
